import os, re, math
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.style import WD_STYLE_TYPE
from docx.text.paragraph import Paragraph
from docx.table import Table

from PyQt5.QtWidgets import (QApplication, QMainWindow, QPushButton, QVBoxLayout, QHBoxLayout, 
                            QFileDialog, QLabel, QWidget, QMessageBox, QGroupBox, 
                            QRadioButton, QButtonGroup, QInputDialog)
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QIcon, QFont, QPalette, QColor, QPixmap
from docx.oxml import OxmlElement
from docx.oxml.ns import qn



# تحديد مسار 
if getattr(sys, 'frozen', False):
    base_path = sys._MEIPASS       # ملفات مضمنة (pe.docx, norm.docx, fonts...)
    app_path = os.path.dirname(sys.executable)  # مكان exe
else:
    base_path = os.path.dirname(os.path.abspath(__file__))  # ملفات القوالب لو شغال بالكود
    app_path = base_path

    # 📂 مجلد resources
resources_path = os.path.join(app_path, "resources")

# 📝 ملفات القوالب والخطوط
pe_template   = os.path.join(resources_path, "pe.docx")
norm_template = os.path.join(resources_path, "norm.docx")
font_regular  = os.path.join(resources_path, "Dubai-Regular.ttf")
font_bold     = os.path.join(resources_path, "Dubai-Bold.ttf")
logo_path     = os.path.join(resources_path, "logo.png")
important     = os.path.join(resources_path, "important.txt")

# تعريف الألوان
BEIGE_COLOR = "#ff9e00"
DARK_BG_COLOR = "#2D2D30"
DARK_TEXT_COLOR = "#F5F5F5"
LIGHT_BG_COLOR = "#F0F0F0"
LIGHT_TEXT_COLOR = "#000000"
LIGHT_BEIGE = "#ff9d0b"
DARK_BEIGE = "#BF7600"
WHITE_COLOR = "#FFFFFF"

# إنشاء مجلدات إذا لم تكن موجودة
text_files_dir = os.path.join(app_path, "Text Files")
word_files_dir = os.path.join(app_path, "Word Files")
if not os.path.exists(text_files_dir):
    os.makedirs(text_files_dir)
if not os.path.exists(word_files_dir):
    os.makedirs(word_files_dir)


def keep_paragraph(paragraph):
    """يُجبر الفقرة أن تبقى مع العنصر التالي (keepNext)."""
    p = paragraph._element
    pPr = p.get_or_add_pPr()
    keepNext = OxmlElement('w:keepNext')
    pPr.append(keepNext)

def set_cant_split_row(row):
    """يضيف <w:cantSplit/> إلى trPr لِمنع تقسيم الصف عبر صفحات/أعمدة."""
    tr = row._tr
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = OxmlElement('w:trPr')
        tr.insert(0, trPr)
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)

def set_paragraph_num_id(paragraph, num_id, ilvl=0):
    """ربط الفقرة برقم تسلسل معين (numId) لتطبيق إعادة الترقيم."""
    pPr = paragraph._element.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr')
        pPr.append(numPr)
    else:
        numPr.clear()
        
    ilvl_elem = OxmlElement('w:ilvl')
    ilvl_elem.set(qn('w:val'), str(ilvl))
    numPr.append(ilvl_elem)
    
    numId_elem = OxmlElement('w:numId')
    numId_elem.set(qn('w:val'), str(num_id))
    numPr.append(numId_elem)

def get_abstract_num_id_for_style(document, style_name="N"):
    """البحث عن abstractNumId الخاص بـ style معينة (مثل N أو M3WAN ANSWER TABLE NUMBERS) في القالب."""
    num_part = document.part.numbering_part
    if num_part is None:
        return None
    target_style = style_name.replace(" ", "").upper()
    for abs_num in num_part._element.findall(qn('w:abstractNum')):
        p_style = abs_num.find('.//' + qn('w:pStyle'))
        if p_style is not None:
            val = p_style.get(qn('w:val'), '').replace(" ", "").upper()
            if val == target_style:
                return abs_num.get(qn('w:abstractNumId'))
    return '3' if style_name == 'N' else '0'

def create_new_num_instance(document, abstract_num_id):
    """إنشاء numId جديد في numbering.xml مبني على abstractNumId لإعادة بدء الترقيم من 1 تلقائياً."""
    num_part = document.part.numbering_part
    if num_part is None:
        return None
    numbering_elem = num_part._element
    
    max_num_id = 0
    for num in numbering_elem.findall(qn('w:num')):
        nid = int(num.get(qn('w:numId'), '0'))
        if nid > max_num_id:
            max_num_id = nid
            
    new_num_id = max_num_id + 1
    num_elem = OxmlElement('w:num')
    num_elem.set(qn('w:numId'), str(new_num_id))
    abs_elem = OxmlElement('w:abstractNumId')
    abs_elem.set(qn('w:val'), str(abstract_num_id))
    num_elem.append(abs_elem)
    numbering_elem.append(num_elem)
    return new_num_id





class M3wanOCRApp(QMainWindow):
    def apply_numbering_style(self, paragraph, document):
        """تطبيق ستايل N مع ضبط numId الحالي لإعادة الترقيم من 1"""
        paragraph.style = document.styles["N"]
        if getattr(self, "current_num_id", None):
            set_paragraph_num_id(paragraph, self.current_num_id)

    def __init__(self):
        super().__init__()
        self.setWindowTitle("M3wanOCRv3.0")
        self.setGeometry(100, 100, 830, 250)
        self.setLayoutDirection(Qt.RightToLeft)
        
        self.load_logo()
        self.dark_mode = True
        self.setup_ui()
        
        self.apply_theme(dark_mode=True)

    def load_logo(self):
        self.logo_path = logo_path
        if os.path.exists(logo_path):
            self.setWindowIcon(QIcon(logo_path))
        
    def setup_ui(self):
        main_widget = QWidget()
        main_layout = QVBoxLayout()
        main_layout.setAlignment(Qt.AlignCenter)
        
        title_layout = QHBoxLayout()
        title_layout.setAlignment(Qt.AlignCenter)
        
        self.logo_path = logo_path
        title_label = QLabel("M3wanOCRv3.0")
        title_font = QFont("Dubai", 18, QFont.Bold)
        title_label.setFont(title_font)
        title_layout.addWidget(title_label)
        
        if os.path.exists(logo_path):
            logo_label = QLabel()
            pixmap = QPixmap(logo_path)
            logo_label.setPixmap(pixmap.scaled(50, 50, Qt.KeepAspectRatio, Qt.SmoothTransformation))
            title_layout.addWidget(logo_label)
        
        theme_group = QGroupBox("السمة")
        theme_group.setAlignment(Qt.AlignCenter)
        theme_layout = QHBoxLayout()
        theme_layout.setAlignment(Qt.AlignCenter)
        
        self.light_radio = QRadioButton("فاتح")
        self.dark_radio = QRadioButton("داكن")
        self.dark_radio.setChecked(True)
        
        theme_button_group = QButtonGroup(self)
        theme_button_group.addButton(self.light_radio)
        theme_button_group.addButton(self.dark_radio)
        
        theme_layout.addWidget(self.light_radio)
        theme_layout.addWidget(self.dark_radio)
        theme_group.setLayout(theme_layout)
        
        self.light_radio.toggled.connect(lambda checked: checked and self.apply_theme(False))
        self.dark_radio.toggled.connect(lambda checked: checked and self.apply_theme(True))
        
        button_layout = QHBoxLayout()
        button_layout.setAlignment(Qt.AlignCenter)
        
        self.edit_button = QPushButton("تعديل")
        self.edit_button.setFixedWidth(130)
        self.edit_button.clicked.connect(self.edit_text)
        
        self.word_to_txt_button = QPushButton("وورد إلى TXT")
        self.word_to_txt_button.setFixedWidth(130)
        self.word_to_txt_button.clicked.connect(self.convert_word_to_txt)

        self.print_button = QPushButton("تحويل إلى وورد")
        self.print_button.setFixedWidth(130)
        self.print_button.clicked.connect(self.print_to_word)

        self.previous_button = QPushButton("سنين سابقة")
        self.previous_button.setFixedWidth(130)
        self.previous_button.clicked.connect(self.print_previous_years)

        self.chart_button = QPushButton("رسم بياني")
        self.chart_button.setFixedWidth(130)
        self.chart_button.clicked.connect(self.show_exam_chart)

        button_layout.addWidget(self.edit_button)
        button_layout.addWidget(self.word_to_txt_button)
        button_layout.addWidget(self.print_button)
        button_layout.addWidget(self.previous_button)
        button_layout.addWidget(self.chart_button)
        
        main_layout.addLayout(title_layout)
        main_layout.addWidget(theme_group)
        main_layout.addLayout(button_layout)
        
        main_widget.setLayout(main_layout)
        self.setCentralWidget(main_widget)
        
        font = QFont("Dubai", 10)
        self.setFont(font)
    
    def apply_theme(self, dark_mode):
        self.dark_mode = dark_mode
        palette = QPalette()
        
        self.theme_label_color = WHITE_COLOR if dark_mode else DARK_BEIGE
        
        if dark_mode:
            palette.setColor(QPalette.Window, QColor(DARK_BG_COLOR))
            palette.setColor(QPalette.WindowText, QColor(LIGHT_BEIGE))
            palette.setColor(QPalette.Base, QColor(DARK_BG_COLOR).darker(120))
            palette.setColor(QPalette.AlternateBase, QColor(DARK_BG_COLOR))
            palette.setColor(QPalette.ToolTipBase, QColor(DARK_BG_COLOR))
            palette.setColor(QPalette.ToolTipText, QColor(LIGHT_BEIGE))
            palette.setColor(QPalette.Text, QColor(LIGHT_BEIGE))
            palette.setColor(QPalette.Button, QColor(DARK_BG_COLOR).darker(110))
            palette.setColor(QPalette.ButtonText, QColor(LIGHT_BEIGE))
            palette.setColor(QPalette.BrightText, Qt.red)
            palette.setColor(QPalette.Link, QColor(BEIGE_COLOR))
            palette.setColor(QPalette.Highlight, QColor(BEIGE_COLOR))
            palette.setColor(QPalette.HighlightedText, QColor(DARK_BG_COLOR))
        else:
            palette.setColor(QPalette.Window, QColor(LIGHT_BG_COLOR))
            palette.setColor(QPalette.WindowText, QColor(DARK_BEIGE))
            palette.setColor(QPalette.Base, Qt.white)
            palette.setColor(QPalette.AlternateBase, QColor(LIGHT_BG_COLOR))
            palette.setColor(QPalette.ToolTipBase, Qt.white)
            palette.setColor(QPalette.ToolTipText, QColor(DARK_BEIGE))
            palette.setColor(QPalette.Text, QColor(DARK_BEIGE))
            palette.setColor(QPalette.Button, QColor(LIGHT_BG_COLOR))
            palette.setColor(QPalette.ButtonText, QColor(DARK_BEIGE))
            palette.setColor(QPalette.BrightText, Qt.red)
            palette.setColor(QPalette.Link, QColor(BEIGE_COLOR))
            palette.setColor(QPalette.Highlight, QColor(BEIGE_COLOR))
            palette.setColor(QPalette.HighlightedText, Qt.white)
        
        QApplication.setPalette(palette)
        self.update_styles()
    
    def update_styles(self):
        button_style = f"""
            QPushButton {{ 
                background-color: {BEIGE_COLOR}; 
                color: black; 
                border: none; 
                border-radius: 4px; 
                padding: 8px;
                font-family: 'Dubai';
            }}
            QPushButton:hover {{ 
                background-color: {'#cc7e00' if self.dark_mode else '#ffb84d'};
            }}
            QPushButton:disabled {{ 
                background-color: {'#666666' if self.dark_mode else '#cccccc'};
                color: {'#aaaaaa' if self.dark_mode else '#888888'};
            }}
        """
        
        group_style = f"""
            QGroupBox {{
                font-weight: bold;
                border: 1px solid {'#555555' if self.dark_mode else '#bbbbbb'};
                border-radius: 5px;
                margin-top: 10px;
                padding-top: 15px;
            }}
            QGroupBox::title {{
                subcontrol-origin: margin;
                subcontrol-position: top center;
                padding: 0 5px;
                color: {WHITE_COLOR if self.dark_mode else DARK_BEIGE};
            }}
        """
        
        self.edit_button.setStyleSheet(button_style)
        self.word_to_txt_button.setStyleSheet(button_style)
        self.print_button.setStyleSheet(button_style)
        self.previous_button.setStyleSheet(button_style)
        self.chart_button.setStyleSheet(button_style)
        
        for group in self.findChildren(QGroupBox):
            group.setStyleSheet(group_style)
    
    def word_length(self, word):
        j = 11
        for i in word:
            if i in [";",":","'",",","."]:
                j=j+2
                continue
            if i in ["i","j","l","I","!"\
                    ,"|"," "]:
                j=j+3
                continue
            if i in ["f","`","\""]:
                j=j+4
                continue
            if i in ["c","r","s","t","z"\
                    ,"-","(",")","{","}"\
                    ,"[","]"]:
                j=j+5
                continue
            if i in ["a","b","d","e","g"\
                    ,"h","k","n","o","p"\
                    ,"q","u","v","x","y"\
                    ,"F","J","L","*","/"\
                    ,"\\","?"]:
                j=j+6
                continue
            if i in ["B","C","E","K","P"\
                    ,"R","S","T","X","Y"\
                    ,"Z","1","2","3","4"\
                    ,"5","6","7","8","9"\
                    ,"0","_","+","=","~"\
                    ,"#","$","^","<",">"]:
                j=j+7
                continue
            if i in ["w","A","D","G","U"\
                    ,"V","&"]:
                j=j+8
                continue
            if i in ["m","H","N","O","Q"]:
                j=j+9
                continue
            if i in ["M"]:
                j=j+10
                continue
            if i in ["W","@"]:
                j=j+11
                continue
            else:
                j=j+12
        return j

    def remove(self, no):
        if not no or not str(no).strip():
            return ""
        no = str(no).strip()
        no_clean = re.sub(r"^(\d+|[a-zA-Z])[\)\-\. ]*", "", no).strip()
        return no_clean.replace("\x0c", "")
    
    def highlight(self, text, words):
        for i in words:
        
            positions = []

            start = 0
            while True:
                start = str(text).lower().find(i, start)
                if start == -1:
                    break
                positions.append(start)
                start += len(i)
            
            print(positions)
            for j in positions:
                print("something")
                text.paragraphs[j:j+len(i)+1].font.highlight_color = WD_COLOR_INDEX.YELLOW
            # start = 0
            # while True:
                # start = str(text).lower().find(i, start)
                # if start == -1:
                    # break
                # text.font.highlight_color = WD_COLOR_INDEX.YELLOW
                # start = start + 1
                
    
    def extract_source_from_text(self, text):
        if not text:
            return text, ""
        match = re.search(r"(\s*(\[[^\]]+\]|\([^\)]+\))\s*)+$", text.strip())
        if match:
            main_text = text[:match.start()].rstrip()
            source_text = text[match.start():].strip()
            return main_text, source_text
        return text, ""

    def save_sources(self):
        if hasattr(self, 'collected_sources') and self.collected_sources:
            unique_sources = list(dict.fromkeys(self.collected_sources))
            sources_path = os.path.join(text_files_dir, "sources.txt")
            try:
                with open(sources_path, "w", encoding="utf-8") as f:
                    for src in unique_sources:
                        f.write(src + "\n")
            except Exception as e:
                print(f"Error saving sources.txt: {e}")

    def question(self, document, y, x, q1, table):
        q_text = self.remove(q1)
        main_text, source_text = self.extract_source_from_text(q_text)
        cell_p = table.cell(y, x).paragraphs[0]
        
        if source_text:
            if not hasattr(self, 'collected_sources'):
                self.collected_sources = []
            self.collected_sources.append(source_text)
            
            run_main = cell_p.add_run(main_text + " ")
            run_main.style = document.styles["Q"]
            
            run_src = cell_p.add_run(source_text)
            run_src.style = document.styles["Q"]
            run_src.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            run = cell_p.add_run(q_text)
            run.style = document.styles["Q"]


    def lettering(self, document, y, x, letter, find_1, find_2, choice, ans, table):
        run = table.cell(y,x).paragraphs[0].add_run(letter)
        if ((choice.startswith("=")) or (ans.lower().find(find_1) > -1)) == True: 
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        table.cell(y,x).paragraphs[0].style = document.styles["L"]

    def choice(self, document, y, x, choice, table):
        run = table.cell(y,x).paragraphs[0].add_run(self.remove(choice))
        run.style = document.styles["M"]
    
    def return_letters_only(self, string):
        return string.replace("0","").replace("1","").replace("2","")\
        .replace("3","").replace("4","").replace("5","").replace("6","")\
        .replace("7","").replace("8","").replace("9","").replace("-","")
    
    def table_0_1(self, document, q1, a, index, ans):
        table = document.add_table(2,2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        run = table.cell(1,1).paragraphs[0].add_run("………………………………………………………………………………………………………………………………………………………………………………….\n………………………………………………………………………………………………………………………………………………………………………………….")
        run.style = document.styles["M"]
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_1_1(self, document, q1, a, index, ans):
        table = document.add_table(2,2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_2_2(self, document, q1, a, b, index, ans):
        table = document.add_table(2,4)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(246)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(246)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(246)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(246)
        table.cell(0,1).merge(table.cell(0,3))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_2_1(self, document, q1, a, b, index, ans):
        table = document.add_table(3,2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,2,0,"B-","b","2",b,ans,table)
        self.choice(document,2,1,b,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)
                    
    def table_3_3(self, document, q1, a, b, c, index, ans):
        table = document.add_table(2, 6)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(150)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(150)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(150)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(150)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(150)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(150)
        table.cell(0,1).merge(table.cell(0,5))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_3_2(self, document, q1, a, b, c, index, ans):
        table = document.add_table(3, 4)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(246)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(246)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(246)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(246)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(246)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(246)
        table.cell(0,1).merge(table.cell(0,3))
        table.cell(2,1).merge(table.cell(2,3))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,2,0,"C-","c","3",c,ans,table)
        self.choice(document,2,1,c,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_3_1(self, document, q1, a, b, c, index, ans):
        table = document.add_table(4, 2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(534)
        table.cell(3,0).width = Pt(42)
        table.cell(3,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,2,0,"B-","b","2",b,ans,table)
        self.choice(document,2,1,b,table)
        self.lettering(document,3,0,"C-","c","3",c,ans,table)
        self.choice(document,3,1,c,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_4_4(self, document, q1, a, b, c, d, index, ans):
        table = document.add_table(2,8)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(102)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(102)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(102)
        table.cell(0,6).width = Pt(42)
        table.cell(0,7).width = Pt(102)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(102)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(102)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(102)
        table.cell(1,6).width = Pt(42)
        table.cell(1,7).width = Pt(102)
        table.cell(0,1).merge(table.cell(0,7))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        self.lettering(document,1,6,"D-","d","4",d,ans,table)
        self.choice(document,1,7,d,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_4_2(self, document, q1, a, b, c, d, index, ans):
        table = document.add_table(3,4)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(246)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(246)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(246)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(246)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(246)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(246)
        table.cell(0,1).merge(table.cell(0,3))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,2,0,"C-","c","3",c,ans,table)
        self.choice(document,2,1,c,table)
        self.lettering(document,2,2,"D-","d","4",d,ans,table)
        self.choice(document,2,3,d,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_4_1(self, document, q1, a, b, c, d, index, ans):
        table = document.add_table(5,2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(534)
        table.cell(3,0).width = Pt(42)
        table.cell(3,1).width = Pt(534)
        table.cell(4,0).width = Pt(42)
        table.cell(4,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,2,0,"B-","b","2",b,ans,table)
        self.choice(document,2,1,b,table)
        self.lettering(document,3,0,"C-","c","3",c,ans,table)
        self.choice(document,3,1,c,table)
        self.lettering(document,4,0,"D-","d","4",d,ans,table)
        self.choice(document,4,1,d,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_5_5(self, document, q1, a, b, c, d, e, index, ans):
        table = document.add_table(2,10)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(73.2)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(73.2)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(73.2)
        table.cell(0,6).width = Pt(42)
        table.cell(0,7).width = Pt(73.2)
        table.cell(0,8).width = Pt(42)
        table.cell(0,9).width = Pt(73.2)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(73.2)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(73.2)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(73.2)
        table.cell(1,6).width = Pt(42)
        table.cell(1,7).width = Pt(73.2)
        table.cell(1,8).width = Pt(42)
        table.cell(1,9).width = Pt(73.2)
        table.cell(0,1).merge(table.cell(0,9))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        self.lettering(document,1,6,"D-","d","4",d,ans,table)
        self.choice(document,1,7,d,table)
        self.lettering(document,1,8,"E-","e","5",e,ans,table)
        self.choice(document,1,9,e,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_5_3(self, document, q1, a, b, c, d, e, index, ans):
        table = document.add_table(3,6)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(150)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(150)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(150)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(150)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(150)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(150)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(150)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(150)
        table.cell(2,4).width = Pt(42)
        table.cell(2,5).width = Pt(150)
        table.cell(0,1).merge(table.cell(0,5))
        table.cell(2,3).merge(table.cell(2,5))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        self.lettering(document,2,0,"D-","d","4",d,ans,table)
        self.choice(document,2,1,d,table)
        self.lettering(document,2,2,"E-","e","5",e,ans,table)
        self.choice(document,2,3,e,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_5_2(self, document, q1, a, b, c, d, e, index, ans):
        table = document.add_table(4,4)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(246)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(246)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(246)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(246)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(246)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(246)
        table.cell(3,0).width = Pt(42)
        table.cell(3,1).width = Pt(246)
        table.cell(3,2).width = Pt(42)
        table.cell(3,3).width = Pt(246)
        table.cell(0,1).merge(table.cell(0,3))
        table.cell(3,1).merge(table.cell(3,3))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,2,0,"C-","c","3",c,ans,table)
        self.choice(document,2,1,c,table)
        self.lettering(document,2,2,"D-","d","4",d,ans,table)
        self.choice(document,2,3,d,table)
        self.lettering(document,3,0,"E-","e","5",e,ans,table)
        self.choice(document,3,1,e,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_5_1(self, document, q1, a, b, c, d, e, index, ans):
        table = document.add_table(6,2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(534)
        table.cell(3,0).width = Pt(42)
        table.cell(3,1).width = Pt(534)
        table.cell(4,0).width = Pt(42)
        table.cell(4,1).width = Pt(534)
        table.cell(5,0).width = Pt(42)
        table.cell(5,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,2,0,"B-","b","2",b,ans,table)
        self.choice(document,2,1,b,table)
        self.lettering(document,3,0,"C-","c","3",c,ans,table)
        self.choice(document,3,1,c,table)
        self.lettering(document,4,0,"D-","d","4",d,ans,table)
        self.choice(document,4,1,d,table)
        self.lettering(document,5,0,"E-","e","5",e,ans,table)
        self.choice(document,5,1,e,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_6_6(self, document, q1, a, b, c, d, e, f, index, ans):
        table = document.add_table(2,12)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(54)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(54)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(54)
        table.cell(0,6).width = Pt(42)
        table.cell(0,7).width = Pt(54)
        table.cell(0,8).width = Pt(42)
        table.cell(0,9).width = Pt(54)
        table.cell(0,10).width = Pt(42)
        table.cell(0,11).width = Pt(54)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(54)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(54)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(54)
        table.cell(1,6).width = Pt(42)
        table.cell(1,7).width = Pt(54)
        table.cell(1,8).width = Pt(42)
        table.cell(1,9).width = Pt(54)
        table.cell(1,10).width = Pt(42)
        table.cell(1,11).width = Pt(54)
        table.cell(0,1).merge(table.cell(0,11))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        self.lettering(document,1,6,"D-","d","4",d,ans,table)
        self.choice(document,1,7,d,table)
        self.lettering(document,1,8,"E-","e","5",e,ans,table)
        self.choice(document,1,9,e,table)
        self.lettering(document,1,10,"F-","f","6",f,ans,table)
        self.choice(document,1,11,f,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_6_4(self, document, q1, a, b, c, d, e, f, index, ans):
        table = document.add_table(3,8)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(102)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(102)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(102)
        table.cell(0,6).width = Pt(42)
        table.cell(0,7).width = Pt(102)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(102)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(102)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(102)
        table.cell(1,6).width = Pt(42)
        table.cell(1,7).width = Pt(102)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(102)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(102)
        table.cell(2,4).width = Pt(42)
        table.cell(2,5).width = Pt(102)
        table.cell(2,6).width = Pt(42)
        table.cell(2,7).width = Pt(102)
        table.cell(0,1).merge(table.cell(0,7))
        table.cell(2,1).merge(table.cell(2,3))
        table.cell(2,5).merge(table.cell(2,7))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        self.lettering(document,1,6,"D-","d","4",d,ans,table)
        self.choice(document,1,7,d,table)
        self.lettering(document,2,0,"E-","e","5",e,ans,table)
        self.choice(document,2,1,e,table)
        self.lettering(document,2,4,"F-","f","6",f,ans,table)
        self.choice(document,2,5,f,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_6_3(self, document, q1, a, b, c, d, e, f, index, ans):
        table = document.add_table(3, 6)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(150)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(150)
        table.cell(0,4).width = Pt(42)
        table.cell(0,5).width = Pt(150)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(150)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(150)
        table.cell(1,4).width = Pt(42)
        table.cell(1,5).width = Pt(150)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(150)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(150)
        table.cell(2,4).width = Pt(42)
        table.cell(2,5).width = Pt(150)
        table.cell(0,1).merge(table.cell(0,5))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,1,4,"C-","c","3",c,ans,table)
        self.choice(document,1,5,c,table)
        self.lettering(document,2,0,"D-","d","4",d,ans,table)
        self.choice(document,2,1,d,table)
        self.lettering(document,2,2,"E-","e","5",e,ans,table)
        self.choice(document,2,3,e,table)
        self.lettering(document,2,4,"F-","f","6",f,ans,table)
        self.choice(document,2,5,f,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_6_2(self, document, q1, a, b, c, d, e, f, index, ans):
        table = document.add_table(4,4)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(246)
        table.cell(0,2).width = Pt(42)
        table.cell(0,3).width = Pt(246)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(246)
        table.cell(1,2).width = Pt(42)
        table.cell(1,3).width = Pt(246)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(246)
        table.cell(2,2).width = Pt(42)
        table.cell(2,3).width = Pt(246)
        table.cell(3,0).width = Pt(42)
        table.cell(3,1).width = Pt(246)
        table.cell(3,2).width = Pt(42)
        table.cell(3,3).width = Pt(246)
        table.cell(0,1).merge(table.cell(0,3))
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,1,2,"B-","b","2",b,ans,table)
        self.choice(document,1,3,b,table)
        self.lettering(document,2,0,"C-","c","3",c,ans,table)
        self.choice(document,2,1,c,table)
        self.lettering(document,2,2,"D-","d","4",d,ans,table)
        self.choice(document,2,3,d,table)
        self.lettering(document,3,0,"E-","e","5",e,ans,table)
        self.choice(document,3,1,e,table)
        self.lettering(document,3,2,"F-","f","6",f,ans,table)
        self.choice(document,3,3,f,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def table_6_1(self, document, q1, a, b, c, d, e, f, index, ans):
        table = document.add_table(7,2)
        table.style = 'T'
        table.cell(0,0).width = Pt(42)
        table.cell(0,1).width = Pt(534)
        table.cell(1,0).width = Pt(42)
        table.cell(1,1).width = Pt(534)
        table.cell(2,0).width = Pt(42)
        table.cell(2,1).width = Pt(534)
        table.cell(3,0).width = Pt(42)
        table.cell(3,1).width = Pt(534)
        table.cell(4,0).width = Pt(42)
        table.cell(4,1).width = Pt(534)
        table.cell(5,0).width = Pt(42)
        table.cell(5,1).width = Pt(534)
        table.cell(6,0).width = Pt(42)
        table.cell(6,1).width = Pt(534)
        self.apply_numbering_style(table.cell(0,0).paragraphs[0], document)
        self.question(document,0,1,q1,table)
        self.lettering(document,1,0,"A-","a","1",a,ans,table)
        self.choice(document,1,1,a,table)
        self.lettering(document,2,0,"B-","b","2",b,ans,table)
        self.choice(document,2,1,b,table)
        self.lettering(document,3,0,"C-","c","3",c,ans,table)
        self.choice(document,3,1,c,table)
        self.lettering(document,4,0,"D-","d","4",d,ans,table)
        self.choice(document,4,1,d,table)
        self.lettering(document,5,0,"E-","e","5",e,ans,table)
        self.choice(document,5,1,e,table)
        self.lettering(document,6,0,"F-","f","6",f,ans,table)
        self.choice(document,6,1,f,table)
        # for row in table.rows:
            # set_cant_split_row(row)
        # for row in table.rows:
            # for cell in row.cells:
                # for paragraph in cell.paragraphs:
                    # keep_paragraph(paragraph)

    def answer_table(self, document,title, a, b, ans):
        answer_list = ans[a:b]
    
        if len(answer_list) == 0: return
    
        table = document.add_table(math.ceil(len(answer_list)/10) * 2 + 1, 10)
        table.cell(0,0).merge(table.cell(0,9))
        table.cell(0,0).paragraphs[0].add_run("Answers - " + title)
        table.style = "M3WAN ANSWER TABLE"

        # إعادة بدء ترقيم جداول الإجابات من 1 دائماً عند كل جدول إجابات جديد
        abs_id_ans = get_abstract_num_id_for_style(document, "M3WAN ANSWER TABLE NUMBERS")
        ans_num_id = create_new_num_instance(document, abs_id_ans)
     
        for i in range(math.ceil(len(answer_list)/10)*10):
            cell_p = table.cell(math.floor(i/10) * 2 + 1, i % 10).paragraphs[0]
            cell_p.style = document.styles["M3WAN ANSWER TABLE NUMBERS"]
            if ans_num_id:
                set_paragraph_num_id(cell_p, ans_num_id)
        
            try:    
                table.cell(math.floor(i/10) * 2 + 2, i % 10).paragraphs[0].add_run(\
                self.return_letters_only(answer_list[i].upper()).replace("#","-"))
            except:
                table.cell(math.floor(i/10) * 2 + 2, i % 10).paragraphs[0].add_run(\
                "-")
    
        document.add_paragraph("‍ ")


    def answer_list(self, x):
        y = ""
        z = ""
        index = 1
        counter = 1
        
        x = x.lower()
        x = x.replace("\n"," ").replace("\r"," ")
        x = x.replace("true","a").replace("false","b")
        
        if len(x) == 0:
            return ""

        for j in x:
            if j in ["a","b","c","d","e","f","/","#","+"]:
                y = y + j
                    
                    
        for j in ["a","b","c","d","e","f","/","#"]:
            y = y.replace(j,j+" ")
        
        
        y = y.replace(" +","")
        
        y = y.split()
        
        
        
        for j in y:
            if j == "/": 
                if counter > 10 : index = 1
                counter = 1
                continue
            z = z + str(index) + "-" + j + " "
            index = index + 1
            counter = counter + 1
       
        
        return z

    def edit_1(self, text_file_1, text_file_2, ans):
        text_1 = open(text_file_1,"r", encoding = "utf-8")
        text_2 = open(text_file_2,"w", encoding = "utf-8")
        line = text_1.readline()
        while line != "":
            if line.strip()[:2] == "//":
                ans = ans + line
                if (line.strip()[-2:] == "//") & (len(line.strip()) > 3):
                    line = text_1.readline()
                    continue
                else:
                    line = text_1.readline()
                    ans = ans + line
                    while line.strip()[-2:] != "//":
                        line = text_1.readline()
                        ans = ans + line
            else: text_2.write(line)
            line = text_1.readline()
        text_1.close()
        text_2.close()
        return ans

    def edit_2(self, text_file_1, text_file_2, ans):
        text_1 = open(text_file_1,"r", encoding = "utf-8")
        text_2 = open(text_file_2,"w", encoding = "utf-8")
        text_2.write(ans.strip())
        for i in text_1:
            i = i.replace("\n","").replace("\r","").strip()
            if (i == "" or i.isdigit() or len(i) == 1):
                continue
            else:
                try:
                    while i[1].isdigit(): i = i[1:]
                except:
                    pass
                text_2.write("\n" + i)
        text_1.close()
        text_2.close()

    def edit_3(self, text_file_1, text_file_2, my_list):
        text_1 = open(text_file_1,"r", encoding = "utf-8")
        text_2 = open(text_file_2,"w", encoding = "utf-8")
        for n,i in enumerate(text_1):
            if n == 0: 
                text_2.write(i)
                continue
            for j in my_list:
                i = i.replace(j,"\n" + j[1:])
            text_2.write(i)
        text_1.close()
        text_2.close()

    def edit_4(self, text_file_1, text_file_2, my_list):
        text_1 = open(text_file_1,"r", encoding = "utf-8")
        text_2 = open(text_file_2,"w", encoding = "utf-8")
        line_1 = text_1.readline().strip()
        while True:
            line_2 = text_1.readline().strip()
            if (line_2[0:2] or line_2[0:3]) in my_list:
                text_2.write(line_1 + "\n")
                line_1 = line_2
            else:
                line_1 = line_1 + " " + line_2
            if line_2 == "":
                text_2.write(line_1)
                break
        text_1.close()
        text_2.close()

    def edit_5(self, text_file_1, text_file_2, my_list):
        text_1 = open(text_file_1,"r", encoding = "utf-8")
        text_2 = open(text_file_2,"w", encoding = "utf-8")
        line = text_1.readline()
        text_2.write(line)
        while True:
            line = text_1.readline()
            if line == "":break
            if line[0:2] in my_list:
                text_2.write(line[0:2] + " " + line[2:])
                continue
            if line[0:3] in my_list:
                text_2.write(line[0:3] + " " + line[3:])
            else: text_2.write(line)
        text_1.close()
        text_2.close()

    def edit_all(self, input_file, output_file):
        ans = ""
        my_list = []
        for i in [" "," (","	","	("]:
            for j in ["a","A","b","B","c","C","d","D","e","E","f","F","g","G"]:
                for k in ["-","#",")",".",",",":",";"]:
                    my_list.append(i+j+k)
        my_list = my_list + [" b "," c "," d "," e "," f "," g "," B "," C "," D "," E "," F "," G "]
        ans = self.answer_list(self.edit_1(input_file,"temp_1.txt",ans))
        self.edit_2("temp_1.txt","temp_2.txt",ans)
        self.edit_3("temp_2.txt","temp_3.txt",my_list)
        new_list = []
        for i in my_list:
            new_list = new_list + i.split()
        my_list = new_list
        for i in range(12):
            my_list.pop(-2)
        for i in ["","("]:
            for j in ["1","2","3","4","5","6","7","8","9","0"]:
                for k in ["-","#",")",".",",",":",";"]:
                    my_list.append(i+j+k)
        my_list.append("1_")
        my_list.append("2_")
        self.edit_4("temp_3.txt","temp_4.txt",my_list)
        my_list.pop(-1)
        my_list.pop(-1)
        self.edit_5("temp_4.txt",output_file,my_list)
        os.remove("temp_1.txt")
        os.remove("temp_2.txt")
        os.remove("temp_3.txt")
        os.remove("temp_4.txt")

    def print_all(self, input_file, output_file):
        # تحديد مسار ملف norm.docx
        template_path = norm_template
        
        # التحقق من وجود الملف
        if not os.path.exists(template_path):
            raise FileNotFoundError("ملف norm.docx غير موجود في المسار المحدد.")
        
        # فتح norm.docx كقالب
        document = Document(template_path)
        abs_id = get_abstract_num_id_for_style(document, "N")
        self.current_num_id = create_new_num_instance(document, abs_id)

        self.collected_sources = []

        
        #فتح ملف الكلمات المهمة
        temp_file = open(important)
        important_words = []
        for i in temp_file:
            important_words.append(i.strip())
        temp_file.close()
        # إضافة المحتوى الجديد كما في الكود الأصلي
        line = ""
        index = 0
        q1 = ""
        q2 = ""
        a = ""
        b = ""
        c = ""
        d = ""
        e = ""
        f = ""
        abcdef = ""
        ans = ""
        
        start_of_table_index = 0
        title_text = ""
        
        text = open(input_file, "r", encoding="utf-8")
        abcdef = text.readline().split()
        q2 = text.readline().strip()
        
        while True:
            if (index % 100) == 0 : document.save(output_file)
            if q2 == "1end":
                p = document.add_paragraph()
                run = p.add_run("‍ ")
                run.font.size = Pt(8)
                run.font.name = "Dubai"
               
                self.answer_table(document,\
                title_text,start_of_table_index,index,abcdef) 
                break
            if q2.strip().startswith("1_"):
                if index != 0: # جدول الإجابات
                    p = document.add_paragraph()
                    run = p.add_run("‍ ")
                    run.font.size = Pt(8)
                    run.font.name = "Dubai"
                    
                    self.answer_table(document,\
                    title_text,start_of_table_index,index,abcdef)
                    
                    start_of_table_index = index
                    document.add_page_break()
                        
                title_text = q2[2:].strip() # عنوان أساسي
                paragraph = document.add_paragraph(title_text)
                paragraph.style = "1-العناوين الأساسية"
                self.current_num_id = create_new_num_instance(document, get_abstract_num_id_for_style(document, "N"))
                
                q2 = text.readline().strip()
                if q2 == "": q2 ="1end"
                continue
            elif q2.strip().startswith("2_"):
                    
                p = document.add_paragraph()
                run = p.add_run("‍ ")
                run.font.size = Pt(8)
                run.font.name = "Dubai"
                
                q2 = q2[2:].strip()
                paragraph = document.add_paragraph(q2)
                paragraph.style = "2- العناوين الفرعية"
                # keep_paragraph(paragraph)

                q2 = text.readline().strip()
                if q2 == "": q2 ="1end"
                continue
            p = document.add_paragraph()
            run = p.add_run("‍ ")
            run.font.size = Pt(8)
            run.font.name = "Dubai"
            try:
                ans = abcdef[index]
            except:
                ans = ""
            index = index + 1
            q1 = q2
            line = text.readline().strip()
            if line == "": line ="1end"
            if line[0].isdigit() == True:
                q2 = line
                self.table_0_1(document,q1,"a-",index,ans)
                continue
            a = line
            line = text.readline().strip()
            if line == "": line = "1end"
            if line[0].isdigit() == True:
                q2 = line
                self.table_1_1(document,q1,a,index,ans)
                continue
            b = line
            line = text.readline().strip()
            if line == "": line = "1end"
            if line[0].isdigit() == True:
                q2 = line
                if max(self.word_length(self.remove(a)),self.word_length(self.remove(b))) <= 246:
                    self.table_2_2(document,q1,a,b,index,ans)
                else:
                    self.table_2_1(document,q1,a,b,index,ans)
                continue
            c = line
            line = text.readline().strip()
            if line == "": line = "1end"
            if line[0].isdigit() == True:
                q2 = line
                if max(self.word_length(self.remove(a)),self.word_length(self.remove(b)), self.word_length(self.remove(c))) <= 150:
                    self.table_3_3(document,q1,a,b,c,index,ans)
                elif max(self.word_length(self.remove(a)),self.word_length(self.remove(b)), self.word_length(self.remove(c))) <= 246:
                    self.table_3_2(document,q1,a,b,c,index,ans)
                else:
                    self.table_3_1(document,q1,a,b,c,index,ans)
                continue
            d = line
            line = text.readline().strip()
            if line == "": line = "1end"
            if line[0].isdigit() == True:
                q2 = line
                if max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d))) <= 102:
                    self.table_4_4(document,q1,a,b,c,d,index,ans)
                elif max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d))) <= 246:
                    self.table_4_2(document,q1,a,b,c,d,index,ans)
                else:
                    self.table_4_1(document,q1,a,b,c,d,index,ans)
                continue
            e = line
            line = text.readline().strip()
            if line == "": line = "1end"
            if line[0].isdigit() == True:
                q2 = line
                if max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e))) <= 73.2:
                    self.table_5_5(document,q1,a,b,c,d,e,index,ans)
                elif max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e))) <= 150:
                    self.table_5_3(document,q1,a,b,c,d,e,index,ans)
                elif max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e))) <= 246:
                    self.table_5_2(document,q1,a,b,c,d,e,index,ans)
                else:
                    self.table_5_1(document,q1,a,b,c,d,e,index,ans)
                continue
            f = line
            line = text.readline().strip()
            if line == "" : line = "1end"
            if line[0].isdigit() == True:
                q2 = line
                if max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e)),self.word_length(self.remove(f))) <= 54:
                    self.table_6_6(document,q1,a,b,c,d,e,f,index,ans)
                elif max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e)),self.word_length(self.remove(f))) <= 102:
                    self.table_6_4(document,q1,a,b,c,d,e,f,index,ans)
                elif max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e)),self.word_length(self.remove(f))) <= 150:
                    self.table_6_3(document,q1,a,b,c,d,e,f,index,ans)
                elif max(self.word_length(self.remove(a)), self.word_length(self.remove(b)), self.word_length(self.remove(c)),self.word_length(self.remove(d)), self.word_length(self.remove(e)),self.word_length(self.remove(f))) <= 246:
                    self.table_6_2(document,q1,a,b,c,d,e,f,index,ans)
                else:
                    self.table_6_1(document,q1,a,b,c,d,e,f,index,ans)
                continue
       
         # فاصل قبل جدول الإجابات
        # document.add_paragraph()

        # ---------------------------
        # جدول الإجابات في النهاية
        # ---------------------------
        # total = len(abcdef)
        # if total > 0:
            # cols = 10
            # rows = (total // cols) * 2 + (2 if total % cols else 0)
            # table = document.add_table(rows=rows, cols=cols)
            # table.style = 'Table Grid'

            # for i in range(total):
                # col = i % cols
                # block = i // cols

                # الصف الأول: أرقام
                # cell_num = table.cell(block * 2, col)
                # cell_num.text = str(i + 1)
                # for p in cell_num.paragraphs:
                    # if p.runs:
                        # run = p.runs[0]
                        # run.font.name = "Dubai"
                        # run.font.size = Pt(11)
                        # run.bold = True
                    # p.alignment = 1  # وسط

                # الصف الثاني: الإجابة
                # cell_ans = table.cell(block * 2 + 1, col)
                # cell_ans.text = self.return_letters_only(abcdef[i]).upper()
                # for p in cell_ans.paragraphs:
                    # if p.runs:
                        # run = p.runs[0]
                        # run.font.name = "Dubai"
                        # run.font.size = Pt(11)
                        # run.bold = True
                    # p.alignment = 1  # وسط

            # document.add_paragraph()  # مسافة بعد الجدول

        document.save(output_file)
        self.save_sources()
        text.close()


    def edit_text(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "اختر ملف TXT", "", "Text Files (*.txt)")
        if file_path:
            try:
                output_path = os.path.join(text_files_dir, "output.txt")
                self.edit_all(file_path, output_path)
                msg = QMessageBox(self)
                msg.setIcon(QMessageBox.Information)
                msg.setWindowTitle("نجاح")
                msg.setText(f"تم تعديل الملف وحفظه في:\n{output_path}")
                msg.setStyleSheet("QLabel { color: black; }")
                msg.exec_()
                os.startfile(os.path.dirname(output_path))
            except Exception as e:
                msg = QMessageBox(self)
                msg.setIcon(QMessageBox.Critical)
                msg.setWindowTitle("خطأ")
                msg.setText(f"حدث خطأ أثناء التعديل:\n{str(e)}")
                msg.setStyleSheet("QLabel { color: black; }")
                msg.exec_()

    def print_to_word(self):
        input_path = os.path.join(text_files_dir, "output.txt")
        if not os.path.exists(input_path):
            msg = QMessageBox(self)
            msg.setIcon(QMessageBox.Warning)
            msg.setWindowTitle("تحذير")
            msg.setText("ملف output.txt غير موجود. قم بإنشائه أولاً باستخدام خيار التعديل.")
            msg.setStyleSheet("QLabel { color: black; }")
            msg.exec_()
            return
        
        # فتح نافذة لإدخال اسم الملف
        file_name, ok = QInputDialog.getText(self, "اسم الملف", "أدخل اسم ملف Word (بدون الامتداد):")
        if ok and file_name:
            output_file = os.path.join(word_files_dir, f"{file_name}.docx")
            try:
                self.print_all(input_path, output_file)
                msg = QMessageBox(self)
                msg.setIcon(QMessageBox.Information)
                msg.setWindowTitle("نجاح")
                msg.setText(f"تم تحويل الملف إلى Word وحفظه في:\n{output_file}")
                msg.setStyleSheet("QLabel { color: black; }")
                msg.exec_()
                os.startfile(os.path.dirname(output_file))
                os.startfile(output_file)
            except Exception as e:
                msg = QMessageBox(self)
                msg.setIcon(QMessageBox.Critical)
                msg.setWindowTitle("خطأ")
                msg.setText(f"حدث خطأ أثناء التحويل:\n{str(e)}")
                msg.setStyleSheet("QLabel { color: black; }")
                msg.exec_()
        else:
            msg = QMessageBox(self)
            msg.setIcon(QMessageBox.Warning)
            msg.setWindowTitle("تحذير")
            msg.setText("لم يتم إدخال اسم للملف.")
            msg.setStyleSheet("QLabel { color: black; }")
            msg.exec_()

    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_COLOR_INDEX
    from PyQt5.QtWidgets import QMessageBox, QInputDialog
    import os, re

    def print_previous_years(self):
        try:
            # التحقق من وجود ملف output.txt
            input_file = os.path.join(text_files_dir, "output.txt")
            if not os.path.exists(input_file):
                QMessageBox.warning(self, "خطأ", "ملف output.txt غير موجود. قم بإنشائه أولاً باستخدام خيار التعديل.")
                return

            # إنشاء مجلد Word Files في نفس مسار البرنامج
            word_dir = os.path.join(app_path, "Word Files")
            os.makedirs(word_dir, exist_ok=True)
            

            # إدخال اسم الملف
            file_name, ok = QInputDialog.getText(self, "تسمية الملف", "ادخل اسم ملف Word (بدون الامتداد):")
            if not ok or not file_name.strip():
                QMessageBox.warning(self, "تحذير", "لم يتم إدخال اسم للملف.")
                return

            save_path = os.path.join(word_dir, f"{file_name.strip()}.docx")

            # التحقق من وجود القالب pe.docx
            template_path = pe_template
            if not os.path.exists(template_path):
                QMessageBox.critical(self, "خطأ", "ملف pe.docx غير موجود في المسار.")
                return

            # فتح القالب
            document = Document(template_path)
            abs_id = get_abstract_num_id_for_style(document, "N")
            self.current_num_id = create_new_num_instance(document, abs_id)

            self.collected_sources = []

            
            #فتح ملف الكلمات المهمة
            temp_file = open(important)
            important_words = []
            for i in temp_file:
                important_words.append(i.strip())
            temp_file.close()
        
            # قراءة النصوص من output.txt
            f = open(input_file,"r", encoding = "utf-8")
            
            # أول سطر = الإجابات
            answers = f.readline().split()
            lines = []
            
            while True:
                line = f.readline().strip()
                if line == "": break
                lines.append(line)
                
            # with open(input_file, "r", encoding="utf-8") as f:
                # lines = [line.strip() for line in f if line.strip()]

            if not lines:
                QMessageBox.warning(self, "خطأ", "ملف output.txt فارغ.")
                return

            # answers = lines[0].split()
            # lines = lines[1:]
            
            start_of_table_index = 0
            title_text = ""
            q_index = -1
            self._choice_counter = 0
            self._last_question_index = -1

            for line in lines:
                if line.strip().startswith("1_"): 
                    if q_index != -1: # جدول الإجابات
                        self.answer_table(document,\
                        title_text,start_of_table_index,q_index+1,answers)
                        
                        start_of_table_index = q_index+1
                        document.add_page_break()
                        
                    title_text = line[2:].strip() # عنوان أساسي
                    paragraph = document.add_paragraph(title_text)
                    paragraph.style = "1-العناوين الأساسية"
                    self.current_num_id = create_new_num_instance(document, get_abstract_num_id_for_style(document, "N"))
                
                
                elif line.strip().startswith("2_"):  # عنوان فرعي
                    secondary_title_text = line[2:].strip()
                    paragraph = document.add_paragraph(secondary_title_text)
                    paragraph.style = "2- العناوين الفرعية"
                    
                    # -------- سؤال --------
                elif re.match(r"^\d+", line):  # يبدأ برقم (لكن مش 2_)
                    q_index += 1
                    clean_line = re.sub(r"^\d+[\)\-\. ]*", "", line).strip()
                    self._choice_counter = 0

                    p = document.add_paragraph()
                    self.apply_numbering_style(p, document)
                    main_text, source_text = self.extract_source_from_text(clean_line)
                    if source_text:
                        if not hasattr(self, 'collected_sources'):
                            self.collected_sources = []
                        self.collected_sources.append(source_text)
                        
                        run1 = p.add_run(main_text + " ")
                        run1.bold = True
                        run1.underline = True
                        run1.font.name = "Dubai"
                        run1.font.size = Pt(12)
                        
                        run2 = p.add_run(source_text)
                        run2.bold = True
                        run2.underline = True
                        run2.font.name = "Dubai"
                        run2.font.size = Pt(12)
                        run2.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        run = p.add_run(clean_line)
                        run.bold = True
                        run.underline = True
                        run.font.name = "Dubai"
                        run.font.size = Pt(12)

                # اختيار
                elif re.match(r"^[a-zA-Z][\)\-\. ]+", line):
                    letter = chr(65 + self._choice_counter)  # 65 = A
                    self._choice_counter += 1

                    option_text = re.sub(r"^[a-zA-Z][\)\-\. ]*", "", line).strip()
                    clean_choice = f"{letter}- {option_text}"

                    p = document.add_paragraph()
                    run = p.add_run(clean_choice)
                    run.font.name = "Dubai"
                    run.font.size = Pt(12)

                    # تمييز الإجابة الصحيحة
                    if q_index < len(answers):
                        correct = answers[q_index].upper()
                        if correct.find(letter) > -1 :
                            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

                # نص إضافي
                else:
                    p = document.add_paragraph()
                    run = p.add_run(line)
                    run.font.name = "Dubai"
                    run.font.size = Pt(12)

            # جدول الإجابات في النهاية
            
            self.answer_table(document,\
            title_text,start_of_table_index,q_index+1,answers)
            
            
            # total = len(answers)
            # if total > 0:
                # document.add_paragraph()
                # cols = 10
                # rows = (total // cols) * 2 + (2 if total % cols else 0)
                # table = document.add_table(rows=rows, cols=cols)
                # table.style = 'Table Grid'
                # for i in range(total):
                    # col = i % cols
                    # block = i // cols
                    # cell_num = table.cell(block * 2, col)
                    # cell_num.text = str(i + 1)
                    # for p in cell_num.paragraphs:
                        # run = p.runs[0]
                        # run.font.name = "Dubai"
                        # run.font.size = Pt(11)
                        # run.bold = True
                        # p.alignment = 1
                    # cell_ans = table.cell(block * 2 + 1, col)
                    # cell_ans.text = self.return_letters_only(answers[i]).upper()
                    # for p in cell_ans.paragraphs:
                        # run = p.runs[0]
                        # run.font.name = "Dubai"
                        # run.font.size = Pt(11)
                        # run.bold = True
                        # p.alignment = 1

            # حفظ الملف
            document.save(save_path)
            self.save_sources()
            QMessageBox.information(self, "تم", f"تم إنشاء الملف:\n{save_path}")

            os.startfile(save_path)
            os.startfile(word_dir)

        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"حدث خطأ:\n{str(e)}")

    def show_exam_chart(self):
        try:
            default_input = os.path.join(text_files_dir, "output.txt")
            if os.path.exists(default_input):
                file_path = default_input
            else:
                file_path, _ = QFileDialog.getOpenFileName(self, "اختر ملف TXT للتحليل", text_files_dir, "Text Files (*.txt)")

            if not file_path or not os.path.exists(file_path):
                QMessageBox.warning(self, "تحذير", "لم يتم اختيار ملف أو الملف غير موجود.")
                return

            with open(file_path, "r", encoding="utf-8") as f:
                answers_raw = f.readline().split()
                lines = [line.strip() for line in f if line.strip()]

            if not lines:
                QMessageBox.warning(self, "خطأ", "الملف فارغ.")
                return

            exam_data = []
            current_exam = None
            q_count = 0

            for line in lines:
                if line.strip().startswith("1_"):
                    if current_exam is not None:
                        exam_data.append({"title": current_exam, "q_count": q_count})
                    current_exam = line[2:].strip()
                    q_count = 0
                elif re.match(r"^\d+", line) and not line.startswith("2_"):
                    q_count += 1

            if current_exam is not None:
                exam_data.append({"title": current_exam, "q_count": q_count})

            if not exam_data:
                QMessageBox.warning(self, "خطأ", "لم يتم العثور على عناوين امتحانات تبدأ بـ 1_ في الملف.")
                return

            total_answers = len(answers_raw)
            total_questions = sum(item["q_count"] for item in exam_data)

            ans_idx = 0
            for item in exam_data:
                qc = item["q_count"]
                allocated = max(0, min(qc, total_answers - ans_idx))
                item["ans_count"] = allocated
                ans_idx += qc

            import matplotlib.pyplot as plt
            import numpy as np

            exam_titles = [f"{i+1}. {item['title']}" for i, item in enumerate(exam_data)]
            q_counts = [item["q_count"] for item in exam_data]
            a_counts = [item["ans_count"] for item in exam_data]

            fig, ax = plt.subplots(figsize=(max(12, len(exam_data) * 0.8), 6), dpi=150)
            x = np.arange(len(exam_titles))
            width = 0.38

            rects1 = ax.bar(x - width/2, q_counts, width, label="questions (الأسئلة)", color="#1f77b4", edgecolor="#134b73")
            rects2 = ax.bar(x + width/2, a_counts, width, label="answers (الإجابات)", color="#ff7f0e", edgecolor="#b35400")

            ax.set_ylabel("Count (العدد)", fontsize=12, fontweight="bold")
            ax.set_title("Questions vs Answers Chart (مقارنة عدد الأسئلة والإجابات لكل امتحان)", fontsize=14, fontweight="bold", pad=15)
            ax.set_xticks(x)

            short_titles = [t[:18] + ("..." if len(t) > 18 else "") for t in exam_titles]
            ax.set_xticklabels(short_titles, rotation=45, ha="right", fontsize=9, fontweight="bold")

            ax.legend(loc="upper right", frameon=True, fontsize=11)
            ax.grid(axis="y", linestyle="--", alpha=0.6)

            for rect in rects1:
                h = rect.get_height()
                ax.annotate(f"{int(h)}", xy=(rect.get_x() + rect.get_width() / 2, h),
                            xytext=(0, 3), textcoords="offset points", ha="center", va="bottom", fontsize=9, fontweight="bold", color="#134b73")

            for rect in rects2:
                h = rect.get_height()
                ax.annotate(f"{int(h)}", xy=(rect.get_x() + rect.get_width() / 2, h),
                            xytext=(0, 3), textcoords="offset points", ha="center", va="bottom", fontsize=9, fontweight="bold", color="#b35400")

            plt.tight_layout()
            charts_dir = os.path.join(app_path, "Charts")
            os.makedirs(charts_dir, exist_ok=True)
            chart_path = os.path.join(charts_dir, "questions_vs_answers_chart.png")
            plt.savefig(chart_path, dpi=150)
            plt.close()

            summary_msg = f"مقارنة الأسئلة والإجابات:\n"
            summary_msg += f"- إجمالي الأسئلة: {total_questions}\n"
            summary_msg += f"- إجمالي الإجابات: {total_answers}\n\n"
            
            mismatch_count = 0
            for item in exam_data:
                diff = item["q_count"] - item["ans_count"]
                if diff == 0:
                    status = "✅ متطابق"
                else:
                    status = f"⚠️ غير متطابق (فارق {diff})"
                    mismatch_count += 1
                summary_msg += f"• {item['title']}: أسئلة={item['q_count']} | إجابات={item['ans_count']} -> {status}\n"

            msg_box = QMessageBox(self)
            msg_box.setIcon(QMessageBox.Warning if mismatch_count > 0 or total_questions != total_answers else QMessageBox.Information)
            msg_box.setWindowTitle("تقرير الرسم البياني")
            msg_box.setText(summary_msg)
            msg_box.setStyleSheet("QLabel { color: black; }")
            msg_box.exec_()

            os.startfile(charts_dir)
            os.startfile(chart_path)

        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء إنشاء الرسم البياني:\n{str(e)}")

    def docx_to_txt(self, docx_path, output_txt_path):
        document = Document(docx_path)

        def is_arabic(text):
            if re.search(r'[a-zA-Z]', text):
                return False
            return bool(re.search(r'[\u0600-\u06FF]', text))

        def is_run_or_element_highlighted(element):
            if hasattr(element, 'font') and element.font.highlight_color:
                return True
            xml_elem = element._element if hasattr(element, '_element') else element
            highlights = xml_elem.xpath('.//w:highlight')
            for h in highlights:
                val = h.get(qn('w:val'))
                if val and val.lower() != 'none':
                    return True
            return False

        # Collect all body children
        body_children = list(document.element.body)

        # Count page breaks in document
        total_page_breaks = 0
        for child in body_children:
            xml_str = child.xml
            if 'w:type="page"' in xml_str or ('w:br' in xml_str and 'type="page"' in xml_str) or 'w:lastRenderedPageBreak' in xml_str:
                total_page_breaks += 1

        current_page = 1
        skip_pages = (total_page_breaks >= 2)

        global_q_counter = 0
        topic_q_counter = 0
        has_title_emitted = False

        section_questions_count = 0
        section_highlight_answers = []
        section_table_answers = []

        table_answers = []
        highlight_answers = []
        txt_content_lines = []

        doc_base_title = os.path.splitext(os.path.basename(docx_path))[0]

        def flush_section():
            nonlocal section_questions_count, section_highlight_answers, section_table_answers
            if section_questions_count == 0 and not section_table_answers:
                return
            
            if len(section_table_answers) > section_questions_count:
                aligned_table = section_table_answers[:section_questions_count]
            elif len(section_table_answers) < section_questions_count:
                aligned_table = section_table_answers + ["-"] * (section_questions_count - len(section_table_answers))
            else:
                aligned_table = section_table_answers

            table_answers.extend(aligned_table)
            highlight_answers.extend(section_highlight_answers)

            section_questions_count = 0
            section_highlight_answers = []
            section_table_answers = []

        for idx, child in enumerate(body_children):
            xml_str = child.xml
            has_break = ('w:type="page"' in xml_str or ('w:br' in xml_str and 'type="page"' in xml_str) or 'w:lastRenderedPageBreak' in xml_str)

            if skip_pages and current_page < 3:
                if has_break:
                    current_page += 1
                continue

            if has_break:
                current_page += 1

            if child.tag.endswith('p'):
                p = Paragraph(child, document)
                raw_text = p.text.replace('\u200d', '').replace('\ufeff', '').strip()
                if not raw_text or is_arabic(raw_text):
                    continue
                p_text = raw_text

                style_name = p.style.name if p.style else ""

                # 1. Main / Sub Heading by Style or starting with 1_ / 2_
                if style_name in ["1-العناوين الأساسية", "2- العناوين الفرعية"] or p_text.startswith("1_") or p_text.startswith("2_"):
                    clean_title = re.sub(r"^[12]_\s*", "", p_text).strip()
                    if clean_title:
                        txt_content_lines.append(f"1_ {clean_title}")
                        topic_q_counter = 0
                        has_title_emitted = True
                    continue

                clean_q = re.sub(r"^\d+[\)\-\. ]*", "", p_text).strip()
                if len(re.findall(r"[a-zA-Z0-9]", clean_q)) < 2:
                    continue

                # 2. Question (starts with digit or style "N")
                if style_name == "N" or re.match(r"^\d+[\)\-\. ]+", p_text):
                    if not has_title_emitted:
                        txt_content_lines.append(f"1_ {doc_base_title}")
                        topic_q_counter = 0
                        has_title_emitted = True

                    global_q_counter += 1
                    topic_q_counter += 1
                    section_questions_count += 1
                    txt_content_lines.append(f"{topic_q_counter}- {clean_q}")
                    section_highlight_answers.append("-")
                    continue

                # 3. Choice (starts with letter e.g. A-, B-)
                choice_match = re.match(r"^([a-zA-Z])[\)\-\. ]+(.*)", p_text)
                if choice_match:
                    letter = choice_match.group(1).upper()
                    choice_val = choice_match.group(2).strip()
                    txt_content_lines.append(f"{letter}- {choice_val}")

                    if is_run_or_element_highlighted(p):
                        if section_highlight_answers and section_questions_count <= len(section_highlight_answers):
                            section_highlight_answers[section_questions_count - 1] = letter
                    continue

                # 4. Other English text: peek at next non-empty non-Arabic paragraph
                next_is_choice = False
                for next_idx in range(idx + 1, len(body_children)):
                    next_child = body_children[next_idx]
                    if next_child.tag.endswith('p'):
                        np_text = Paragraph(next_child, document).text.replace('\u200d', '').replace('\ufeff', '').strip()
                        if np_text and not is_arabic(np_text):
                            if re.match(r"^([a-zA-Z])[\)\-\. ]+", np_text):
                                next_is_choice = True
                            break

                if next_is_choice:
                    if not has_title_emitted:
                        txt_content_lines.append(f"1_ {doc_base_title}")
                        topic_q_counter = 0
                        has_title_emitted = True

                    global_q_counter += 1
                    topic_q_counter += 1
                    section_questions_count += 1
                    txt_content_lines.append(f"{topic_q_counter}- {clean_q}")
                    section_highlight_answers.append("-")
                else:
                    clean_title = p_text.strip()
                    txt_content_lines.append(f"1_ {clean_title}")
                    topic_q_counter = 0
                    has_title_emitted = True

            elif child.tag.endswith('tbl'):
                table = Table(child, document)
                style_name = table.style.name if table.style else ""

                first_cell_text = ""
                if len(table.rows) > 0 and len(table.rows[0].cells) > 0:
                    first_cell_text = table.rows[0].cells[0].text.strip()

                if is_arabic(first_cell_text):
                    continue

                # Check if Answer Table
                if style_name == "M3WAN ANSWER TABLE" or "Answers -" in first_cell_text or "Answers" in first_cell_text:
                    for row_idx in range(1, len(table.rows)):
                        if row_idx % 2 == 0:  # rows 2, 4, 6...
                            for cell in table.rows[row_idx].cells:
                                val = cell.text.strip().upper()
                                if val and val != "-" and val != "ANSWERS":
                                    letters = re.findall(r"[A-F]", val)
                                    if letters:
                                        section_table_answers.extend(letters)
                    flush_section()
                    continue

                # Check if Question Table vs TOC Table
                has_choices = False
                for r in table.rows:
                    for c in r.cells:
                        ctext = c.text.strip()
                        if re.match(r"^([A-F])[\-\.\)]*$", ctext, re.IGNORECASE):
                            has_choices = True
                            break
                    if has_choices:
                        break

                if not has_choices:
                    continue

                # Question Table (norm.docx style 'T')
                visited_cells = set()
                q_text = ""
                for row in table.rows:
                    for cell in row.cells:
                        ctext = cell.text.strip()
                        if ctext and not re.match(r"^([A-F])[\-\.\)]*$", ctext, re.IGNORECASE):
                            q_text = ctext
                            break
                    if q_text:
                        break

                if q_text and not is_arabic(q_text):
                    clean_q = re.sub(r"^\d+[\)\-\. ]*", "", q_text).strip()
                    if len(re.findall(r"[a-zA-Z0-9]", clean_q)) < 2:
                        continue

                    if not has_title_emitted:
                        txt_content_lines.append(f"1_ {doc_base_title}")
                        topic_q_counter = 0
                        has_title_emitted = True

                    global_q_counter += 1
                    topic_q_counter += 1
                    section_questions_count += 1
                    txt_content_lines.append(f"{topic_q_counter}- {clean_q}")
                    section_highlight_answers.append("-")

                    current_choice_letter = None
                    for row in table.rows:
                        for cell in row.cells:
                            if cell._tc in visited_cells:
                                continue
                            visited_cells.add(cell._tc)

                            ctext = cell.text.strip()
                            if is_arabic(ctext):
                                continue

                            letter_match = re.match(r"^([A-F])[\-\.\)]*$", ctext, re.IGNORECASE)
                            if letter_match:
                                current_choice_letter = letter_match.group(1).upper()
                                if is_run_or_element_highlighted(cell):
                                    if section_highlight_answers and section_questions_count <= len(section_highlight_answers):
                                        section_highlight_answers[section_questions_count - 1] = current_choice_letter
                                continue

                            if current_choice_letter and ctext:
                                if is_run_or_element_highlighted(cell):
                                    if section_highlight_answers and section_questions_count <= len(section_highlight_answers):
                                        section_highlight_answers[section_questions_count - 1] = current_choice_letter
                                txt_content_lines.append(f"{current_choice_letter}- {ctext}")
                                current_choice_letter = None

        flush_section()

        # Compare and merge answers
        raw_final_answers = []
        mismatches = []
        max_ans_len = max(len(table_answers), len(highlight_answers), global_q_counter)

        for i in range(max_ans_len):
            tbl_ans = table_answers[i] if i < len(table_answers) else "-"
            hl_ans = highlight_answers[i] if i < len(highlight_answers) else "-"

            if tbl_ans != "-" and hl_ans != "-":
                if tbl_ans == hl_ans:
                    raw_final_answers.append(tbl_ans)
                else:
                    mismatches.append(f"السؤال {i+1}: الجدول = {tbl_ans} ، التظليل = {hl_ans}")
                    raw_final_answers.append(tbl_ans)
            elif tbl_ans != "-":
                raw_final_answers.append(tbl_ans)
            elif hl_ans != "-":
                raw_final_answers.append(hl_ans)
            else:
                if i < global_q_counter:
                    raw_final_answers.append("-")

        # Format answers with question numbers e.g. 1-A 2-C 3-# 4-B
        formatted_answers = []
        for i, ans in enumerate(raw_final_answers):
            q_num = i + 1
            ans_char = ans if ans != "-" else "#"
            formatted_answers.append(f"{q_num}-{ans_char}")

        # Write output text file
        with open(output_txt_path, "w", encoding="utf-8") as f:
            f.write(" ".join(formatted_answers) + "\n")
            for line in txt_content_lines:
                f.write(line + "\n")

        return {
            "q_count": global_q_counter,
            "table_ans_count": len(table_answers),
            "highlight_ans_count": len([a for a in highlight_answers if a != "-"]),
            "final_ans_count": len(raw_final_answers),
            "mismatches": mismatches
        }

    def convert_word_to_txt(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "اختر ملف Word", word_files_dir, "Word Files (*.docx)")
        if not file_path:
            return

        default_save_name = os.path.splitext(os.path.basename(file_path))[0] + ".txt"
        default_save_path = os.path.join(text_files_dir, default_save_name)

        save_path, _ = QFileDialog.getSaveFileName(self, "حفظ ملف TXT", default_save_path, "Text Files (*.txt)")
        if not save_path:
            return

        try:
            res = self.docx_to_txt(file_path, save_path)
            
            msg_text = f"تم تحويل ملف الوورد بنجاح وحفظه في:\n{save_path}\n\n"
            msg_text += f"📊 إحصائيات المستند:\n"
            msg_text += f"• عدد الأسئلة المستخرجة: {res['q_count']}\n"
            msg_text += f"• عدد إجابات الجدول: {res['table_ans_count']}\n"
            msg_text += f"• عدد إجابات التظليل الأصفر: {res['highlight_ans_count']}\n"
            msg_text += f"• إجمالي الإجابات النهائي: {res['final_ans_count']}\n"

            if res['mismatches']:
                msg_text += f"\n⚠️ تنبيه: تم العثور على اختلاف بين الجدول والتظليل في {len(res['mismatches'])} سؤال:\n"
                for m in res['mismatches'][:5]:
                    msg_text += f"  - {m}\n"
                if len(res['mismatches']) > 5:
                    msg_text += f"  ...وغيرها ({len(res['mismatches']) - 5} أسئلة أخرى)\n"
            elif res['q_count'] == res['final_ans_count'] and res['q_count'] > 0:
                msg_text += f"\n✅ تطابق تام بين عدد الأسئلة والإجابات!"

            msg = QMessageBox(self)
            msg.setIcon(QMessageBox.Information if not res['mismatches'] else QMessageBox.Warning)
            msg.setWindowTitle("نجاح التحويل")
            msg.setText(msg_text)
            msg.setStyleSheet("QLabel { color: black; }")
            msg.exec_()

            os.startfile(os.path.dirname(save_path))
        except Exception as e:
            msg = QMessageBox(self)
            msg.setIcon(QMessageBox.Critical)
            msg.setWindowTitle("خطأ")
            msg.setText(f"حدث خطأ أثناء تحويل ملف الوورد:\n{str(e)}")
            msg.setStyleSheet("QLabel { color: black; }")
            msg.exec_()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    
    font_path = font_regular
    font_bold_path = font_bold
    
    font_id = QFont("Dubai").family()
    if font_id == "Dubai" or os.path.exists(font_path):
        pass
    else:
        from PyQt5.QtGui import QFontDatabase
        QFontDatabase.addApplicationFont(font_path)
        QFontDatabase.addApplicationFont(font_bold_path)
    
    window = M3wanOCRApp()
    window.show()
    sys.exit(app.exec_())